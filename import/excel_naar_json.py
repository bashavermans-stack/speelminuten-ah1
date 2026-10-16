"""Zet de speelminuten-Excels om naar één JSON-bestand voor de app.

Gebruik:  python import/excel_naar_json.py
Leest:    Data_in_een_programmatje/*.xlsx
          import/aliassen_spelers.csv, import/aliassen_tegenstanders.csv
Schrijft: data/voetbal-data.json
          import/controlerapport.md
"""

from __future__ import annotations

import csv
import datetime as dt
import glob
import json
import os
import re
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass, field

import openpyxl

HIER = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HIER)
BRON = os.path.join(ROOT, "Data_in_een_programmatje")
UIT_JSON = os.path.join(ROOT, "data", "voetbal-data.json")
UIT_RAPPORT = os.path.join(HIER, "controlerapport.md")
ALIAS_SPELERS = os.path.join(HIER, "aliassen_spelers.csv")
ALIAS_TEGENSTANDERS = os.path.join(HIER, "aliassen_tegenstanders.csv")

# Seizoenen waarvan een deel alleen als totaal (zonder wedstrijddetail) bestaat.
ALLEEN_TOTAAL_DEEL = {"2010-2011": "1e helft (alleen totalen bekend)"}

TUSSENVOEGSELS = {"van", "vd", "v", "de", "der", "den", "het", "t", "ter", "ten", "v.d.", "v.d", "vd."}
LABEL_RE = re.compile(r"\((k|[a-z]\d|o\d\d)\)", re.IGNORECASE)


# ---------------------------------------------------------------- namen

def zonder_accenten(s: str) -> str:
    s = unicodedata.normalize("NFKD", s)
    return "".join(ch for ch in s if not unicodedata.combining(ch))


def schoon_naam(voor: str, achter: str) -> tuple[str, str, bool]:
    """Haalt labels als (K) weg en zet tussenvoegsels vóór de achternaam."""
    keeper = bool(re.search(r"\(k\)", f"{voor} {achter}", re.IGNORECASE))
    voor = LABEL_RE.sub("", voor).replace("*", " ")
    achter = LABEL_RE.sub("", achter).replace("*", " ")
    v = voor.split()
    a = achter.split()
    tv: list[str] = []
    # "Jan van | Dijk": tussenvoegsel achter de voornaam
    while len(v) > 1 and v[-1].lower() in TUSSENVOEGSELS:
        tv.insert(0, v.pop())
    # "Dijk van" / "Berg vd": tussenvoegsel achter de achternaam
    achteraan: list[str] = []
    while len(a) > 1 and a[-1].lower() in TUSSENVOEGSELS:
        achteraan.insert(0, a.pop())
    voornaam = " ".join(v)
    achternaam = " ".join(tv + achteraan + a)
    return voornaam, achternaam, keeper


def naam_sleutel(voornaam: str, achternaam: str) -> str:
    def deel(s: str) -> str:
        woorden = [w for w in zonder_accenten(s).lower().split() if w not in TUSSENVOEGSELS]
        return re.sub(r"[^a-z]", "", "".join(woorden))
    return f"{deel(voornaam)}|{deel(achternaam)}"


def slug(s: str) -> str:
    s = zonder_accenten(s).lower()
    return re.sub(r"[^a-z0-9]+", "-", s).strip("-")


def tegenstander_sleutel(s: str) -> str:
    return re.sub(r"[^a-z0-9]", "", zonder_accenten(s).lower())


def levenshtein(a: str, b: str) -> int:
    if len(a) < len(b):
        a, b = b, a
    vorige = list(range(len(b) + 1))
    for i, ca in enumerate(a, 1):
        huidige = [i]
        for j, cb in enumerate(b, 1):
            huidige.append(min(vorige[j] + 1, huidige[j - 1] + 1, vorige[j - 1] + (ca != cb)))
        vorige = huidige
    return vorige[-1]


def lees_alias_csv(pad: str) -> dict[str, str]:
    """CSV met kolommen variant;wordt  (sleutels of namen)."""
    if not os.path.exists(pad):
        return {}
    uit = {}
    with open(pad, encoding="utf-8-sig", newline="") as f:
        for rij in csv.DictReader(f, delimiter=";"):
            variant = (rij.get("variant") or "").strip()
            wordt = (rij.get("wordt") or "").strip()
            if variant and wordt and not variant.startswith("#"):
                uit[variant] = wordt
    return uit


# ---------------------------------------------------------------- Excel lezen

@dataclass
class Blok:
    bestand: str
    sheet: str
    seizoen: str
    kolommen: list[int]
    dagen: dict[int, int]
    maanden: dict[int, int]
    tegenstanders: dict[int, str]
    spelers: list[dict] = field(default_factory=list)  # {voor, achter, totaal, wed, minuten{kol: min}}
    notities: list[dict] = field(default_factory=list)  # {voor, achter, tekst}


def seizoen_uit_tekst(tekst: str) -> str | None:
    m = re.search(r"(20\d\d)\s*-\s*(20\d\d)", tekst)
    return f"{m.group(1)}-{m.group(2)}" if m else None


def is_getal(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def als_int(v) -> int | None:
    """Getal, ook als het als tekst in de cel staat (bijv. '10')."""
    if is_getal(v):
        return int(v)
    if isinstance(v, str) and v.strip().isdigit():
        return int(v.strip())
    return None


def cel(rij: list, i: int):
    return rij[i] if 0 <= i < len(rij) else None


def lees_blokken(pad: str) -> list[Blok]:
    wb = openpyxl.load_workbook(pad, data_only=True, read_only=True)
    blokken: list[Blok] = []
    for ws in wb.worksheets:
        rijen = [list(r) for r in ws.iter_rows(values_only=True)]
        for i, rij in enumerate(rijen):
            j = next((k for k, v in enumerate(rij)
                      if isinstance(v, str) and v.strip().lower() in ("wed", "weds", "wedstr")), None)
            if j is None or i < 2 or j < 2:
                continue
            seizoen = None
            for k in range(max(0, i - 3), i):
                for v in rijen[k]:
                    if isinstance(v, str) and seizoen_uit_tekst(v):
                        seizoen = seizoen_uit_tekst(v)
            if not seizoen:
                continue
            dagen_r, maanden_r = rijen[i - 2], rijen[i - 1]
            kolommen = [c for c in range(j + 1, len(rij))
                        if isinstance(rij[c], str) and rij[c].strip()
                        and als_int(cel(dagen_r, c)) is not None and als_int(cel(maanden_r, c)) is not None]
            blok = Blok(os.path.basename(pad), ws.title, seizoen, kolommen,
                        {c: als_int(dagen_r[c]) for c in kolommen},
                        {c: als_int(maanden_r[c]) for c in kolommen},
                        {c: rij[c].strip() for c in kolommen})
            gestart = False
            einde = len(rijen)
            for r_idx in range(i + 1, len(rijen)):
                p = rijen[r_idx]
                voor, achter = cel(p, j - 2), cel(p, j - 1)
                heeft_naam = isinstance(voor, str) and voor.strip() or isinstance(achter, str) and achter.strip()
                if not heeft_naam:
                    if gestart:
                        einde = r_idx
                        break
                    continue
                gestart = True
                voor = str(voor or "").strip()
                achter = str(achter or "").strip()
                wedcel = cel(p, j)
                if isinstance(wedcel, str):
                    blok.notities.append({"voor": voor, "achter": achter, "tekst": wedcel.strip()})
                    continue
                minuten = {c: int(round(p[c])) for c in kolommen
                           if is_getal(cel(p, c)) and p[c] > 0}
                totaal = cel(p, j + 1)
                blok.spelers.append({"voor": voor, "achter": achter,
                                     "totaal": int(totaal) if is_getal(totaal) else None,
                                     "wed": int(wedcel) if is_getal(wedcel) else None,
                                     "minuten": minuten})
            # notities verderop onder het blok (na de totaalregel)
            leeg = 0
            for p in rijen[einde:]:
                if any(isinstance(v, str) and v.strip().lower() in ("wed", "weds", "wedstr") for v in p):
                    break
                tekst = cel(p, j)
                voor, achter = cel(p, j - 2), cel(p, j - 1)
                if isinstance(tekst, str) and (isinstance(achter, str) or isinstance(voor, str)):
                    blok.notities.append({"voor": str(voor or "").strip(), "achter": str(achter or "").strip(),
                                          "tekst": tekst.strip()})
                    leeg = 0
                elif not any(v not in (None, "") for v in p):
                    leeg += 1
                    if leeg > 4:
                        break
            blokken.append(blok)
    return blokken


def lees_totaal_sheet(pad: str) -> dict[tuple[str, str], dict[str, int]]:
    wb = openpyxl.load_workbook(pad, data_only=True, read_only=True)
    if "totaal" not in wb.sheetnames:
        return {}
    rijen = list(wb["totaal"].iter_rows(values_only=True))
    kop = rijen[0]
    uit = {}
    for r in rijen[1:]:
        if not (isinstance(r[0], str) or isinstance(r[1], str)):
            continue
        per = {}
        for c in range(3, len(kop)):
            m = re.match(r"(\d\d)-(\d\d)", str(kop[c] or ""))
            if m and is_getal(cel(r, c)) and r[c] > 0:
                per[f"20{m.group(1)}-20{m.group(2)}"] = int(r[c])
        uit[(str(r[0] or "").strip(), str(r[1] or "").strip())] = per
    return uit


# ---------------------------------------------------------------- basis/invaller schatten

def schat_basis(minuten: dict[str, int], seizoensminuten: dict[str, int]) -> tuple[dict[str, bool], bool]:
    """Geeft per speler of hij (geschat) in de basis stond, plus of het twijfelachtig was."""
    vol = [s for s, m in minuten.items() if m >= 90]
    rest = sorted((s for s in minuten if minuten[s] < 90),
                  key=lambda s: (-minuten[s], -seizoensminuten.get(s, 0), s))
    k = max(0, 11 - len(vol))
    basis = {s: True for s in vol}
    for idx, s in enumerate(rest):
        basis[s] = idx < k
    twijfel = 0 < k < len(rest) and minuten[rest[k - 1]] == minuten[rest[k]]
    return basis, twijfel


def speler_id(sleutel: str) -> str:
    return "p-" + slug(sleutel.replace("|", " "))


def is_deelreeks(kort: str, lang: str) -> bool:
    """Komen alle letters van kort in volgorde voor in lang? ('bboys' in 'buitenboys')"""
    it = iter(lang)
    return len(kort) >= 3 and all(ch in it for ch in kort)


def koppel_notities(blok, blok_wedstrijden, speler_sleutel, deelname_index, gebeurtenissen, fouten) -> None:
    """Koppelt notities als 'Rood Castricum' of '2x geel TAVV uit' aan de juiste wedstrijd."""
    bekende = {speler_sleutel(p["voor"], p["achter"])[0] for p in blok.spelers}
    for n in blok.notities:
        tekst = zonder_accenten(n["tekst"].lower())
        if "rood" not in tekst and "geel" not in tekst:
            continue
        zonder_x = re.sub(r"\d+\s*x\s*", " ", tekst)
        m_min = re.search(r"(\d+)", zonder_x)
        minuut = int(m_min.group(1)) if m_min else None
        rest = re.sub(r"\b(geel|rood|tegen|uit|thuis|minuut|e)\b|\d+e?", " ", zonder_x)
        t_k = tegenstander_sleutel(rest)
        s = speler_sleutel(n["voor"], n["achter"])[0]
        if s not in bekende:
            dichtbij = sorted(bekende, key=lambda b: levenshtein(b, s))
            if dichtbij and levenshtein(dichtbij[0], s) <= 2:
                s = dichtbij[0]
        kandidaten = [(wid, minuten) for c, (wid, minuten) in blok_wedstrijden.items()
                      if s in minuten and is_deelreeks(tegenstander_sleutel(blok.tegenstanders[c]), t_k)]
        if len(kandidaten) > 1 and minuut is not None:
            kandidaten = [k for k in kandidaten if k[1][s] == minuut] or kandidaten
        if len(kandidaten) > 1:
            kandidaten = [k for k in kandidaten if k[1][s] < 90] or kandidaten
        if len(kandidaten) != 1:
            fouten.append(f"{blok.seizoen}: notitie niet gekoppeld: {n['voor']} {n['achter']} – {n['tekst']}")
            continue
        wid, minuten = kandidaten[0]
        pid = speler_id(s)
        if minuut is None and minuten[s] < 90:
            minuut = minuten[s]
        gebeurtenissen.append({"id": f"g-{wid}-{pid[2:]}-rood", "wedstrijdId": wid, "type": "rood",
                               "minuut": minuut, "spelerId": pid, "opmerking": n["tekst"]})
        deelname_index[(wid, pid)]["gewisseld"] = False


# ---------------------------------------------------------------- hoofdprogramma

def main() -> None:
    bestanden = sorted(glob.glob(os.path.join(BRON, "*.xlsx")))
    rapport: list[str] = ["# Controlerapport import speelminuten", "",
                          f"Gegenereerd: {dt.datetime.now():%Y-%m-%d %H:%M}", ""]

    # 1. Per seizoen de blokken uit het meest recente bestand
    per_seizoen: dict[str, list[Blok]] = {}
    for pad in bestanden:  # oplopend op datum, dus later overschrijft eerder
        gevonden: dict[str, list[Blok]] = defaultdict(list)
        for blok in lees_blokken(pad):
            gevonden[blok.seizoen].append(blok)
        per_seizoen.update(gevonden)
    totaal_sheet: dict = {}
    for pad in bestanden:
        t = lees_totaal_sheet(pad)
        if t:
            totaal_sheet = t

    spelers_alias = lees_alias_csv(ALIAS_SPELERS)          # sleutel -> sleutel
    tegen_alias = lees_alias_csv(ALIAS_TEGENSTANDERS)      # sleutel -> nette naam

    def speler_sleutel(voor: str, achter: str) -> tuple[str, str, str, bool]:
        vn, an, keeper = schoon_naam(voor, achter)
        s = naam_sleutel(vn, an)
        seen = set()
        while s in spelers_alias and s not in seen:
            seen.add(s)
            s = spelers_alias[s]
        return s, vn, an, keeper

    # 2. Spelers verzamelen
    spelers: dict[str, dict] = {}
    varianten: dict[str, Counter] = defaultdict(Counter)
    laatste_seizoen_van: dict[str, str] = {}
    for seizoen in sorted(per_seizoen):
        for blok in per_seizoen[seizoen]:
            for p in blok.spelers:
                s, vn, an, keeper = speler_sleutel(p["voor"], p["achter"])
                sp = spelers.setdefault(s, {"voornaam": vn, "achternaam": an, "keeper": False})
                sp["keeper"] = sp["keeper"] or keeper
                if p["minuten"] or p["totaal"]:
                    # meest recente schrijfwijze wint
                    sp["voornaam"], sp["achternaam"] = vn, an
                    laatste_seizoen_van[s] = seizoen
                varianten[s][f"{p['voor']} | {p['achter']}"] += 1

    # 3. Wedstrijden en deelnames
    wedstrijden, deelnames, gebeurtenissen, seizoenstotalen = [], [], [], []
    deelname_index: dict[tuple[str, str], dict] = {}
    tegenstanders: dict[str, dict] = {}
    tegen_varianten: dict[str, Counter] = defaultdict(Counter)
    fouten: list[str] = []
    twijfel_wedstrijden = 0
    rapport += ["## Seizoenen", "", "| Seizoen | Bestand | Wedstrijden | Spelers |", "|---|---|---|---|"]

    def tegenstander_id(ruw: str) -> str:
        k = tegenstander_sleutel(ruw)
        naam = tegen_alias.get(k, ruw.strip())
        tid = "t-" + slug(naam)
        tegenstanders.setdefault(tid, {"id": tid, "naam": naam, "aliassen": set()})
        tegenstanders[tid]["aliassen"].add(ruw.strip())
        tegen_varianten[tid][ruw.strip()] += 1
        return tid

    for seizoen in sorted(per_seizoen):
        jaar1, jaar2 = (int(x) for x in seizoen.split("-"))
        blokken = per_seizoen[seizoen]
        seizoen_minuten: Counter = Counter()
        verwacht_totaal: dict[str, int] = {}
        verwacht_wed: dict[str, int] = {}
        for blok in blokken:
            for p in blok.spelers:
                s = speler_sleutel(p["voor"], p["achter"])[0]
                for m in p["minuten"].values():
                    seizoen_minuten[s] += m
                if p["totaal"] is not None:
                    verwacht_totaal[s] = max(verwacht_totaal.get(s, 0), p["totaal"])
                if p["wed"] is not None:
                    verwacht_wed[s] = max(verwacht_wed.get(s, 0), p["wed"])

        gezien: set[tuple] = set()
        speler_wed = Counter()
        aantal = 0
        for blok in blokken:
            blok_wedstrijden: dict[int, tuple[str, dict[str, int]]] = {}
            for c in blok.kolommen:
                minuten = {}
                for p in blok.spelers:
                    if c in p["minuten"]:
                        s = speler_sleutel(p["voor"], p["achter"])[0]
                        minuten[s] = minuten.get(s, 0) + p["minuten"][c]
                if not minuten:
                    continue  # toekomstige of lege wedstrijd
                maand, dag = blok.maanden[c], blok.dagen[c]
                jaar = jaar1 if maand >= 7 else jaar2
                try:
                    datum = dt.date(jaar, maand, dag)
                except ValueError:
                    fouten.append(f"{seizoen}: ongeldige datum {dag}-{maand} ({blok.tegenstanders[c]}), overgeslagen")
                    continue
                tid = tegenstander_id(blok.tegenstanders[c])
                if (datum, tid) in gezien:
                    continue
                gezien.add((datum, tid))
                aantal += 1
                wid = f"w-{datum.isoformat()}-{tid[2:]}"
                wedstrijden.append({"id": wid, "seizoen": seizoen, "datum": datum.isoformat(),
                                    "tegenstanderId": tid, "thuisUit": None,
                                    "doelpuntenVoor": None, "doelpuntenTegen": None, "bron": "excel"})
                som = sum(minuten.values())
                if som > 990:
                    fouten.append(f"{datum} {blok.tegenstanders[c]}: totaal {som} min (> 990)")
                if sum(1 for m in minuten.values() if m >= 90) > 11:
                    fouten.append(f"{datum} {blok.tegenstanders[c]}: meer dan 11 spelers met 90 min")
                if any(m > 90 for m in minuten.values()):
                    fouten.append(f"{datum} {blok.tegenstanders[c]}: speler met meer dan 90 min")
                basis, twijfel = schat_basis(minuten, seizoen_minuten)
                twijfel_wedstrijden += twijfel
                for s, m in sorted(minuten.items()):
                    speler_wed[s] += 1
                    d = {"wedstrijdId": wid, "spelerId": speler_id(s), "minuten": min(m, 90),
                         "basis": basis[s], "gewisseld": basis[s] and m < 90, "basisBron": "geschat"}
                    deelnames.append(d)
                    deelname_index[(wid, d["spelerId"])] = d
                blok_wedstrijden[c] = (wid, minuten)
            koppel_notities(blok, blok_wedstrijden, speler_sleutel, deelname_index, gebeurtenissen, fouten)

        # controle seizoenstotalen + aanvullen deel zonder detail
        afwijkingen = 0
        for s, verwacht in verwacht_totaal.items():
            berekend = seizoen_minuten.get(s, 0)
            if verwacht == berekend:
                continue
            if seizoen in ALLEEN_TOTAAL_DEEL and verwacht > berekend:
                seizoenstotalen.append({"seizoen": seizoen, "spelerId": speler_id(s),
                                        "minuten": verwacht - berekend,
                                        "wedstrijden": max(0, verwacht_wed.get(s, 0) - speler_wed.get(s, 0)),
                                        "opmerking": ALLEEN_TOTAAL_DEEL[seizoen]})
                continue
            afwijkingen += 1
            fouten.append(f"{seizoen}: {s} – Excel-totaal {verwacht}, berekend {berekend}")
        rapport.append(f"| {seizoen} | {blokken[0].bestand} | {aantal} | {len(speler_wed)} |"
                       + (f" ⚠ {afwijkingen} afwijkingen" if afwijkingen else ""))

    # controle tegen de totaal-sheet
    for (voor, achter), per in totaal_sheet.items():
        s = speler_sleutel(voor, achter)[0]
        for seizoen, verwacht in per.items():
            berekend = sum(d["minuten"] for d in deelnames
                           if d["spelerId"] == speler_id(s)
                           and d["wedstrijdId"] in {w["id"] for w in wedstrijden if w["seizoen"] == seizoen})
            berekend += sum(t["minuten"] for t in seizoenstotalen
                            if t["spelerId"] == speler_id(s) and t["seizoen"] == seizoen)
            if berekend != verwacht:
                fouten.append(f"totaal-sheet {seizoen}: {voor} {achter} – {verwacht}, berekend {berekend}")

    # 4. Spelerslijst + kandidaten voor samenvoegen
    laatste = max(per_seizoen)
    speler_lijst = []
    for s, sp in sorted(spelers.items()):
        if s not in laatste_seizoen_van:
            continue
        naam = f"{sp['voornaam']} {sp['achternaam']}".strip()
        speler_lijst.append({"id": speler_id(s), "voornaam": sp["voornaam"],
                             "achternaam": sp["achternaam"], "naam": naam, "keeper": sp["keeper"],
                             "actief": laatste_seizoen_van[s] == laatste,
                             "aliassen": sorted(varianten[s])})
    ids = [p["id"] for p in speler_lijst]
    assert len(ids) == len(set(ids)), "dubbele speler-id's"

    kandidaten = []
    sleutels = sorted(laatste_seizoen_van)
    for i, a in enumerate(sleutels):
        va, aa = a.split("|")
        for b in sleutels[i + 1:]:
            vb, ab = b.split("|")
            zelfde_achter = aa == ab and (va[:1] == vb[:1] or levenshtein(va, vb) <= 2)
            zelfde_voor = va == vb and levenshtein(aa, ab) <= 2
            if zelfde_achter or zelfde_voor:
                kandidaten.append((a, b))

    rapport += ["", f"Totaal: {len(wedstrijden)} wedstrijden, {len(speler_lijst)} spelers, "
                    f"{len(tegenstanders)} tegenstanders, {len(deelnames)} deelnames.",
                f"Wedstrijden waarin de basis-schatting een gelijke stand had: {twijfel_wedstrijden}.", ""]
    rapport += ["## Afwijkingen en meldingen", ""] + ([f"- {f}" for f in fouten] or ["- geen"]) + [""]
    rapport += ["## Mogelijke dubbele spelers (controleer!)", "",
                "Samenvoegen: zet in `aliassen_spelers.csv` een regel `variant;wordt` met de sleutels hieronder.", ""]
    rapport += [f"- `{a}` ({', '.join(varianten[a])})  ↔  `{b}` ({', '.join(varianten[b])})"
                for a, b in kandidaten] or ["- geen"]
    rapport += ["", "## Tegenstanders", "", "| Naam | Schrijfwijzen in Excel |", "|---|---|"]
    for tid in sorted(tegenstanders):
        rapport.append(f"| {tegenstanders[tid]['naam']} | {', '.join(sorted(tegen_varianten[tid]))} |")
    rapport += ["", "## Spelers", "", "| Speler | Schrijfwijzen in Excel |", "|---|---|"]
    for p in speler_lijst:
        rapport.append(f"| {p['naam']}{' (K)' if p['keeper'] else ''} | {'; '.join(p['aliassen'])} |")

    data = {
        "versie": 1,
        "gegenereerd": dt.datetime.now().isoformat(timespec="seconds"),
        "bron": "excel-import",
        "spelers": speler_lijst,
        "tegenstanders": [{"id": t["id"], "naam": t["naam"], "aliassen": sorted(t["aliassen"])}
                          for t in sorted(tegenstanders.values(), key=lambda t: t["naam"].lower())],
        "wedstrijden": sorted(wedstrijden, key=lambda w: w["datum"]),
        "deelnames": deelnames,
        "gebeurtenissen": gebeurtenissen,
        "seizoenstotalen": seizoenstotalen,
    }
    os.makedirs(os.path.dirname(UIT_JSON), exist_ok=True)
    with open(UIT_JSON, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=1)
    with open(UIT_RAPPORT, "w", encoding="utf-8") as f:
        f.write("\n".join(rapport) + "\n")
    print(f"{len(wedstrijden)} wedstrijden, {len(speler_lijst)} spelers, {len(tegenstanders)} tegenstanders, "
          f"{len(fouten)} meldingen, {len(kandidaten)} mogelijke dubbelen")
    print(f"-> {UIT_JSON}\n-> {UIT_RAPPORT}")


if __name__ == "__main__":
    main()
